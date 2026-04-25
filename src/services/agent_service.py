from __future__ import annotations

import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from src.agent_core import TravelOrchestrator
from src.agent_core.tool_executor import list_pending_tool_approvals, resolve_tool_approval
from src.agent_core.context_manager import estimate_tokens
from src.config import ASSISTANT_ATTACHMENT_DIR, SUMMARY_MD_DIR, get_context_budget_config
from src.db import get_conn
from src.services.llm_service import LLMService


KEEP_HEAD_MESSAGES = 10
KEEP_TAIL_MESSAGES = 20
MIN_MANUAL_COMPRESS_MESSAGES = 7


def _utc_now() -> str:
  return datetime.now(timezone.utc).isoformat()


def start_new_conversation(user_id: str, title: str | None = None) -> dict[str, str]:
  now = _utc_now()
  conversation_id = str(uuid.uuid4())
  final_title = title or f"New Chat {now[:16].replace('T', ' ')}"
  with get_conn() as conn:
    conn.execute(
      """
      INSERT INTO conversations (id, user_id, title, created_at, updated_at)
      VALUES (?, ?, ?, ?, ?)
      """,
      (conversation_id, user_id, final_title, now, now),
    )
  return {"id": conversation_id, "title": final_title}


def list_conversations(user_id: str, limit: int = 40) -> list[dict[str, Any]]:
  with get_conn() as conn:
    rows = conn.execute(
      """
      SELECT
        c.id,
        c.title,
        c.created_at,
        c.updated_at,
        COUNT(m.id) AS message_count
      FROM conversations c
      LEFT JOIN agent_memory m
        ON c.id = m.conversation_id
        AND m.role IN ('user', 'assistant')
      WHERE c.user_id = ?
      GROUP BY c.id, c.title, c.created_at, c.updated_at
      ORDER BY c.updated_at DESC
      LIMIT ?
      """,
      (user_id, limit),
    ).fetchall()
  return [dict(row) for row in rows]


def ensure_active_conversation(user_id: str) -> str:
  with get_conn() as conn:
    latest = conn.execute(
      """
      SELECT id
      FROM conversations
      WHERE user_id = ?
      ORDER BY updated_at DESC
      LIMIT 1
      """,
      (user_id,),
    ).fetchone()
  if latest:
    return str(latest["id"])

  conv = start_new_conversation(user_id=user_id, title="Default Chat")
  _bind_legacy_messages_to_conversation(user_id=user_id, conversation_id=conv["id"])
  return conv["id"]


def _bind_legacy_messages_to_conversation(user_id: str, conversation_id: str) -> None:
  with get_conn() as conn:
    conn.execute(
      """
      UPDATE agent_memory
      SET conversation_id = ?
      WHERE user_id = ?
        AND role IN ('user', 'assistant')
        AND conversation_id IS NULL
      """,
      (conversation_id, user_id),
    )


def get_chat_history(user_id: str, conversation_id: str | None = None, limit: int = 120) -> list[dict[str, str]]:
  conv_id = conversation_id or ensure_active_conversation(user_id)
  with get_conn() as conn:
    rows = conn.execute(
      """
      SELECT role, content
      FROM (
        SELECT role, content, created_at, id
        FROM agent_memory
        WHERE user_id = ?
          AND conversation_id = ?
          AND role IN ('user', 'assistant')
        ORDER BY created_at DESC, id DESC
        LIMIT ?
      ) t
      ORDER BY created_at ASC, id ASC
      """,
      (user_id, conv_id, limit),
    ).fetchall()
  return [dict(row) for row in rows]


def add_chat_message(
  user_id: str,
  role: str,
  content: str,
  conversation_id: str | None = None,
) -> str:
  conv_id = conversation_id or ensure_active_conversation(user_id)
  now = _utc_now()
  with get_conn() as conn:
    conn.execute(
      """
      INSERT INTO agent_memory (conversation_id, user_id, role, content, created_at)
      VALUES (?, ?, ?, ?, ?)
      """,
      (conv_id, user_id, role, content, now),
    )
    conn.execute("UPDATE conversations SET updated_at = ? WHERE id = ?", (now, conv_id))
  _auto_compress_if_needed(user_id=user_id, conversation_id=conv_id)
  return conv_id


def _auto_compress_if_needed(user_id: str, conversation_id: str) -> None:
  threshold = get_context_budget_config().auto_compress_threshold_tokens
  with get_conn() as conn:
    rows = conn.execute(
      """
      SELECT content
      FROM agent_memory
      WHERE user_id = ?
        AND conversation_id = ?
        AND role IN ('user', 'assistant')
      """,
      (user_id, conversation_id),
    ).fetchall()
  total_tokens = sum(estimate_tokens(str(r["content"])) for r in rows)
  if total_tokens > threshold:
    compress_conversation_history(user_id, conversation_id)


def compress_conversation_history(
  user_id: str,
  conversation_id: str,
  *,
  force: bool = False,
  export_md: bool = False,
) -> dict[str, Any]:
  with get_conn() as conn:
    rows = conn.execute(
      """
      SELECT id, role, content, created_at
      FROM agent_memory
      WHERE user_id = ?
        AND conversation_id = ?
        AND role IN ('user', 'assistant')
      ORDER BY created_at ASC, id ASC
      """,
      (user_id, conversation_id),
    ).fetchall()

    keep_head = KEEP_HEAD_MESSAGES
    keep_tail = KEEP_TAIL_MESSAGES
    if force:
      if len(rows) < MIN_MANUAL_COMPRESS_MESSAGES:
        return {"compressed": False, "message": "Conversation is too short to compress right now."}
      keep_head = max(2, min(KEEP_HEAD_MESSAGES, len(rows) // 4))
      keep_tail = max(3, min(KEEP_TAIL_MESSAGES, len(rows) // 2))
      while keep_head + keep_tail >= len(rows):
        if keep_tail > 3:
          keep_tail -= 1
          continue
        if keep_head > 2:
          keep_head -= 1
          continue
        break
      if keep_head + keep_tail >= len(rows):
        return {"compressed": False, "message": "Conversation is too short to compress right now."}
    else:
      if len(rows) <= KEEP_HEAD_MESSAGES + KEEP_TAIL_MESSAGES:
        return {"compressed": False, "message": "Conversation is too short to compress."}

    middle = rows[keep_head:-keep_tail]
    summary = _summarize_middle_messages([dict(item) for item in middle])
    middle_ids = [int(item["id"]) for item in middle]
    placeholders = ",".join(["?"] * len(middle_ids))
    conn.execute(f"DELETE FROM agent_memory WHERE id IN ({placeholders})", middle_ids)
    conn.execute(
      """
      INSERT INTO agent_memory (conversation_id, user_id, role, content, created_at)
      VALUES (?, ?, 'assistant', ?, ?)
      """,
      (conversation_id, user_id, f"[History Summary]\n{summary}", str(middle[0]["created_at"])),
    )
    conn.execute("UPDATE conversations SET updated_at = ? WHERE id = ?", (_utc_now(), conversation_id))
    md_path = ""
    if export_md:
      md_path = _export_summary_markdown(
        user_id=user_id,
        conversation_id=conversation_id,
        summary_text=summary,
        created_at=str(middle[0]["created_at"]),
        conn=conn,
      )

  return {
    "compressed": True,
    "removed_messages": len(middle),
    "kept_messages": keep_head + keep_tail + 1,
    "mode": "manual" if force else "auto",
    "summary_preview": summary[:220],
    "summary_md_path": md_path,
  }


def get_conversation_compress_hint(user_id: str, conversation_id: str) -> dict[str, Any]:
  threshold = get_context_budget_config().auto_compress_threshold_tokens
  with get_conn() as conn:
    rows = conn.execute(
      """
      SELECT content
      FROM agent_memory
      WHERE user_id = ?
        AND conversation_id = ?
        AND role IN ('user', 'assistant')
      """,
      (user_id, conversation_id),
    ).fetchall()
  total_tokens = sum(estimate_tokens(str(r["content"])) for r in rows)
  ratio = 0.0 if threshold <= 0 else (total_tokens / threshold)
  return {
    "conversation_tokens": total_tokens,
    "threshold_tokens": threshold,
    "ratio": ratio,
    "recommended": ratio >= 0.7,
  }


def _summarize_middle_messages(middle_messages: list[dict[str, str]]) -> str:
  if not middle_messages:
    return "No content to summarize."

  llm = LLMService()
  raw_text = "\n".join([f"{m['role']}: {m['content']}" for m in middle_messages[:80]])
  if llm.is_enabled():
    try:
      return llm.chat(
        system_prompt=(
          "You are a conversation summarizer. Keep key decisions, user intent, places, and constraints. "
          "Return 6-8 concise bullets."
        ),
        messages=[{"role": "user", "content": raw_text}],
      )
    except Exception:
      pass

  bullets: list[str] = []
  for item in middle_messages[:8]:
    text = str(item.get("content", "")).replace("\n", " ").strip()
    if len(text) > 42:
      text = text[:42] + "..."
    prefix = "User" if item.get("role") == "user" else "Assistant"
    bullets.append(f"- {prefix}: {text}")
  if len(middle_messages) > 8:
    bullets.append(f"- Omitted: {len(middle_messages) - 8} more messages")
  return "\n".join(bullets)


def get_memory_notes(user_id: str, limit: int = 100) -> list[dict[str, str]]:
  with get_conn() as conn:
    active_rows = conn.execute(
      """
      SELECT id, content, memory_type, confidence, created_at
      FROM memory_items
      WHERE user_id = ?
        AND is_active = 1
      ORDER BY id DESC
      LIMIT 500
      """,
      (user_id,),
    ).fetchall()
    inactive_rows = conn.execute(
      """
      SELECT content
      FROM memory_items
      WHERE user_id = ?
        AND is_active = 0
      """,
      (user_id,),
    ).fetchall()
    legacy_rows = conn.execute(
      """
      SELECT content, created_at
      FROM agent_memory
      WHERE user_id = ?
        AND role = 'memory'
      ORDER BY id DESC
      LIMIT 500
      """,
      (user_id,),
    ).fetchall()

  inactive_contents = {str(r["content"]).strip() for r in inactive_rows if str(r["content"]).strip()}
  active_items = [dict(r) for r in active_rows]
  active_contents = {str(r.get("content", "")).strip() for r in active_items if str(r.get("content", "")).strip()}

  merged: list[dict[str, str]] = list(active_items)
  for row in legacy_rows:
    content = str(row["content"] or "").strip()
    if not content:
      continue
    if content.startswith("[PENDING_CONFLICT]"):
      continue
    # Hide only conflict-deactivated memories; keep other historical memories.
    if content in inactive_contents:
      continue
    if content in active_contents:
      continue
    merged.append({"content": content, "created_at": str(row["created_at"] or "")})

  merged.sort(key=lambda x: str(x.get("created_at", "")), reverse=True)
  return merged[: max(1, limit)]


def answer(
  user_id: str,
  query: str,
  spots: list[dict[str, Any]],
  conversation_id: str | None = None,
  extra_context: str = "",
) -> str:
  q = query.strip()
  if not q:
    return "请问有什么可以帮助您？"

  orchestrator = TravelOrchestrator()
  history = get_chat_history(user_id=user_id, conversation_id=conversation_id, limit=120)
  ctx_text = (extra_context or "").strip()
  if not ctx_text and conversation_id:
    ctx_text = build_attachment_context(user_id=user_id, conversation_id=conversation_id, limit_chars=2800)

  # Guard: pypdf not installed
  if ctx_text and "Install pypdf" in ctx_text and len(ctx_text) < 300:
    return (
      "检测到您上传了 PDF 文件，但服务器缺少解析库。\n\n"
      "请运行：`pip install pypdf`\n\n"
      "重启服务后即可分析 PDF 内容。"
    )

  reply = orchestrator.run(
    user_id=user_id,
    query=q,
    spots=spots,
    conversation_id=conversation_id,
    history=history,
    extra_context=ctx_text,
  )

  # Fallback: if reply is empty or a bare "Done." placeholder
  _trivial = {"", "Done.", "[MemoryAgent] Done.", "[GeoAgent] Done.",
              "[SocialAgent] Done.", "[PlanAgent] Done."}
  if not reply or reply.strip() in _trivial:
    if ctx_text:
      llm = LLMService()
      if llm.is_enabled():
        try:
          return llm.chat(
            system_prompt=(
              "你是智能旅行助手。根据提供的文档内容和用户问题，给出详细、有条理的分析回复。"
            ),
            messages=[{
              "role": "user",
              "content": f"用户问题：{q}\n\n文档内容：\n{ctx_text}",
            }],
          )
        except Exception:
          pass
      return "已收到您的文件，但暂时无法生成分析结果，请稍后重试或换一种提问方式。"
    return "请问有什么可以帮助您？"

  return reply


def get_agent_runtime_info() -> dict[str, str]:
  llm = LLMService()
  mode = "LLM_TOOL_AGENT" if llm.is_enabled() else "RULE_FALLBACK"
  return {
    "mode": mode,
    "provider": llm.provider_label(),
    "architecture": "orchestrator+geo/social/memory",
  }


def get_latest_agent_debug(user_id: str, conversation_id: str | None = None) -> dict[str, Any]:
  with get_conn() as conn:
    if conversation_id:
      row = conn.execute(
        """
        SELECT *
        FROM agent_run_logs
        WHERE user_id = ? AND conversation_id = ?
        ORDER BY id DESC
        LIMIT 1
        """,
        (user_id, conversation_id),
      ).fetchone()
    else:
      row = conn.execute(
        """
        SELECT *
        FROM agent_run_logs
        WHERE user_id = ?
        ORDER BY id DESC
        LIMIT 1
        """,
        (user_id,),
      ).fetchone()
    if not row:
      return {}
    run = dict(row)
    tools = conn.execute(
      """
      SELECT route_agent, tool_name, tool_args, tool_result, created_at
      FROM agent_tool_logs
      WHERE request_id = ?
      ORDER BY id ASC
      """,
      (run["request_id"],),
    ).fetchall()
    guards = conn.execute(
      """
      SELECT event_type, payload_json, created_at
      FROM agent_guard_logs
      WHERE request_id = ?
      ORDER BY id ASC
      """,
      (run["request_id"],),
    ).fetchall()

  import json

  return {
    "request_id": run["request_id"],
    "route_agent": run["route_agent"],
    "query_text": run["query_text"],
    "memory_hits": json.loads(run["memory_hits_json"] or "[]"),
    "context_budget": json.loads(run["context_budget_json"] or "{}"),
    "history_used": run["history_used"],
    "memory_used": run["memory_used"],
    "tools": [dict(t) for t in tools],
    "guards": [
      {
        "event_type": str(g["event_type"]),
        "payload": json.loads(g["payload_json"] or "{}"),
        "created_at": g["created_at"],
      }
      for g in guards
    ],
    "created_at": run["created_at"],
  }


def get_pending_tool_approvals(user_id: str, limit: int = 30) -> list[dict[str, Any]]:
  return list_pending_tool_approvals(user_id=user_id, limit=limit)


def handle_tool_approval(user_id: str, pending_id: int, action: str) -> dict[str, Any]:
  return resolve_tool_approval(user_id=user_id, pending_id=pending_id, action=action)


def get_request_trace(user_id: str, request_id: str) -> dict[str, Any]:
  rid = request_id.strip()
  if not rid:
    return {}
  with get_conn() as conn:
    run = conn.execute(
      """
      SELECT *
      FROM agent_run_logs
      WHERE user_id = ? AND request_id = ?
      ORDER BY id DESC
      LIMIT 1
      """,
      (user_id, rid),
    ).fetchone()
    if not run:
      return {}
    run_d = dict(run)
    tools = conn.execute(
      """
      SELECT route_agent, tool_name, tool_args, tool_result, created_at
      FROM agent_tool_logs
      WHERE request_id = ?
      ORDER BY id ASC
      """,
      (rid,),
    ).fetchall()
    guards = conn.execute(
      """
      SELECT event_type, payload_json, created_at
      FROM agent_guard_logs
      WHERE request_id = ?
      ORDER BY id ASC
      """,
      (rid,),
    ).fetchall()
  import json

  return {
    "request_id": rid,
    "route_agent": run_d.get("route_agent"),
    "query_text": run_d.get("query_text"),
    "created_at": run_d.get("created_at"),
    "memory_hits": json.loads(run_d.get("memory_hits_json") or "[]"),
    "context_budget": json.loads(run_d.get("context_budget_json") or "{}"),
    "tools": [dict(t) for t in tools],
    "guards": [
      {
        "event_type": str(g["event_type"]),
        "payload": json.loads(g["payload_json"] or "{}"),
        "created_at": g["created_at"],
      }
      for g in guards
    ],
  }


def get_latest_history_summary(user_id: str, conversation_id: str) -> str:
  with get_conn() as conn:
    row = conn.execute(
      """
      SELECT content
      FROM agent_memory
      WHERE user_id = ?
        AND conversation_id = ?
        AND role = 'assistant'
        AND content LIKE '[History Summary]%'
      ORDER BY id DESC
      LIMIT 1
      """,
      (user_id, conversation_id),
    ).fetchone()
  if not row:
    return ""
  return str(row["content"] or "")


def get_latest_history_summary_md(user_id: str, conversation_id: str) -> dict[str, str]:
  with get_conn() as conn:
    row = conn.execute(
      """
      SELECT file_path, summary_text, created_at
      FROM conversation_summaries
      WHERE user_id = ? AND conversation_id = ?
      ORDER BY id DESC
      LIMIT 1
      """,
      (user_id, conversation_id),
    ).fetchone()
  if not row:
    return {}
  return {
    "file_path": str(row["file_path"]),
    "summary_text": str(row["summary_text"] or ""),
    "created_at": str(row["created_at"]),
  }


def save_assistant_attachment(
  *,
  user_id: str,
  conversation_id: str,
  file_name: str,
  mime_type: str,
  data: bytes,
) -> dict[str, Any]:
  safe_name = _safe_file_name(file_name or "attachment.bin")
  now = _utc_now()
  folder = ASSISTANT_ATTACHMENT_DIR / user_id / conversation_id
  folder.mkdir(parents=True, exist_ok=True)
  file_path = folder / f"{uuid.uuid4().hex}_{safe_name}"
  file_path.write_bytes(data)
  extracted = _extract_attachment_text(file_path=file_path, mime_type=mime_type, raw=data)
  with get_conn() as conn:
    conn.execute(
      """
      INSERT INTO assistant_attachments (
        user_id, conversation_id, file_name, mime_type, file_path, extracted_text, created_at
      ) VALUES (?, ?, ?, ?, ?, ?, ?)
      """,
      (
        user_id,
        conversation_id,
        safe_name,
        mime_type or "",
        str(file_path),
        extracted,
        now,
      ),
    )
  return {
    "ok": True,
    "file_name": safe_name,
    "file_path": str(file_path),
    "extracted_preview": extracted[:260],
  }


def list_assistant_attachments(user_id: str, conversation_id: str, limit: int = 20) -> list[dict[str, Any]]:
  with get_conn() as conn:
    rows = conn.execute(
      """
      SELECT id, file_name, mime_type, file_path, extracted_text, created_at
      FROM assistant_attachments
      WHERE user_id = ? AND conversation_id = ?
      ORDER BY id DESC
      LIMIT ?
      """,
      (user_id, conversation_id, limit),
    ).fetchall()
  return [dict(r) for r in rows]


def delete_assistant_attachment(user_id: str, attachment_id: int) -> dict[str, Any]:
  with get_conn() as conn:
    row = conn.execute(
      "SELECT file_path FROM assistant_attachments WHERE id = ? AND user_id = ?",
      (attachment_id, user_id),
    ).fetchone()
    if not row:
      return {"ok": False, "error": "not_found"}
    file_path = str(row["file_path"])
    conn.execute(
      "DELETE FROM assistant_attachments WHERE id = ? AND user_id = ?",
      (attachment_id, user_id),
    )
  try:
    Path(file_path).unlink(missing_ok=True)
  except Exception:
    pass
  return {"ok": True}


def clear_conversation_attachments(user_id: str, conversation_id: str) -> dict[str, Any]:
  with get_conn() as conn:
    rows = conn.execute(
      "SELECT file_path FROM assistant_attachments WHERE user_id = ? AND conversation_id = ?",
      (user_id, conversation_id),
    ).fetchall()
    conn.execute(
      "DELETE FROM assistant_attachments WHERE user_id = ? AND conversation_id = ?",
      (user_id, conversation_id),
    )
  for row in rows:
    try:
      Path(str(row["file_path"])).unlink(missing_ok=True)
    except Exception:
      pass
  return {"ok": True, "deleted": len(rows)}


def build_attachment_context(user_id: str, conversation_id: str, limit_chars: int = 2800) -> str:
  rows = list_assistant_attachments(user_id, conversation_id, limit=8)
  if not rows:
    return ""
  blocks: list[str] = []
  for idx, r in enumerate(rows, start=1):
    text = str(r.get("extracted_text") or "").strip()
    if not text:
      continue
    blocks.append(f"[Attachment {idx}] {r.get('file_name')}\n{text[:900]}")
  merged = "\n\n".join(blocks)
  if len(merged) > limit_chars:
    return merged[:limit_chars]
  return merged


def _export_summary_markdown(
  *,
  user_id: str,
  conversation_id: str,
  summary_text: str,
  created_at: str,
  conn,
) -> str:
  folder = SUMMARY_MD_DIR / user_id / conversation_id
  folder.mkdir(parents=True, exist_ok=True)
  file_path = folder / f"history_summary_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.md"
  content = "# History Summary\n\n" + summary_text.strip() + "\n"
  file_path.write_text(content, encoding="utf-8")
  conn.execute(
    """
    INSERT INTO conversation_summaries (user_id, conversation_id, file_path, summary_text, created_at)
    VALUES (?, ?, ?, ?, ?)
    """,
    (user_id, conversation_id, str(file_path), summary_text, created_at),
  )
  return str(file_path)


def _safe_file_name(name: str) -> str:
  base = os.path.basename(name).strip()
  if not base:
    return "attachment.bin"
  cleaned = "".join(ch for ch in base if ch.isalnum() or ch in {"_", "-", ".", " "}).strip()
  return cleaned or "attachment.bin"


def _extract_attachment_text(*, file_path: Path, mime_type: str, raw: bytes) -> str:
  suffix = file_path.suffix.lower()
  if suffix in {".txt", ".md", ".csv", ".json"}:
    return raw.decode("utf-8", errors="ignore")[:4000]

  if suffix == ".pdf":
    try:
      from pypdf import PdfReader  # type: ignore

      reader = PdfReader(str(file_path))
      text = []
      for page in reader.pages[:8]:
        text.append(page.extract_text() or "")
      return "\n".join(text)[:4000]
    except Exception:
      return f"[{file_path.name}] PDF uploaded. Install pypdf to extract text."

  if suffix in {".docx"}:
    try:
      from docx import Document  # type: ignore

      doc = Document(str(file_path))
      text = "\n".join(p.text for p in doc.paragraphs[:300])
      return text[:4000]
    except Exception:
      return f"[{file_path.name}] Word uploaded. Install python-docx to extract text."

  if suffix in {".xlsx", ".xls"}:
    try:
      import pandas as pd  # type: ignore

      df = pd.read_excel(str(file_path), nrows=120)
      return df.to_csv(index=False)[:4000]
    except Exception:
      return f"[{file_path.name}] Excel uploaded. Install pandas/openpyxl to extract text."

  return f"[{file_path.name}] Attachment uploaded (mime={mime_type or 'unknown'})."
